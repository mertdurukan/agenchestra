# -*- coding: utf-8 -*-
"""Ev stili sade Excel: CSV/JSON -> xlsx (Özet sayfası + en çok 2 veri sayfası).

Kullanım (Windows'ta Türkçe çıktı için -X utf8 önerilir):
  python -X utf8 -I sade_excel.py --cikti <teslim klasörü>/ad_YYYYMMDD.xlsx \
      --baslik "Başlık" [--alt-baslik "tarih · kaynak · kapsam"] \
      --veri a.csv [--sayfa Ad] [--sutunlar "kol1,kol2=Yeni Başlık"] [--aciklama "..."] \
      [--veri b.json --sayfa Ad2 ...] [--ozet ozet.md] [--dondur-sutun 1] [--masaustu]

Girdi: CSV (UTF-8/UTF-8 BOM, ayraç otomatik: , ; | sekme) ya da JSON
(nesne listesi veya {"satirlar": [...]}). --sayfa/--sutunlar/--aciklama, --veri ile
sırayla eşlenir (1. --sayfa -> 1. --veri).
--ozet: düz metin/md; "## BAŞLIK" satırı bölüm başlığı, diğer satırlar madde, boş satır boşluk.
Çıktı geri okunup sayfa ve satır sayıları denetlenir; tutarsızsa çıkış kodu 2.
Sır kalıbı (parola, bağlantı dizesi, anahtar) bulunursa dosya YAZILMAZ, çıkış kodu 3.
"""
import argparse
import csv
import io
import datetime as dt
import json
import os
import re
import shutil
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import TYPE_STRING
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---- ev stili: koyu lacivert başlık, donmuş başlık satırı, filtre, otomatik sütun genişliği ----
KOYU = "1F3864"
GRI = "595959"
BASLIK_FILL = PatternFill("solid", fgColor=KOYU)
BASLIK_FONT = Font(bold=True, color="FFFFFF")
INCE = Side(style="thin", color="BFBFBF")
KENAR = Border(top=INCE, bottom=INCE, left=INCE, right=INCE)
MIN_GEN, MAX_GEN = 8, 60
MAKS_VERI_SAYFASI = 2

SIR_KESIN = [re.compile(p, re.I) for p in (
    r"password\s*[=:]\s*\S", r"pwd\s*=", r"AccountKey\s*=", r"SharedAccessSignature",
    r"Bearer\s+[A-Za-z0-9\-_.]{10,}", r"Server\s*=.*;", r"User\s*Id\s*=",
    r"postgres(ql)?://[^\s:]+:[^\s@]+@", r"-----BEGIN [A-Z ]*PRIVATE KEY-----")]
SIR_SUPHE = re.compile(r"(?<![A-Za-z0-9+/])[A-Za-z0-9+/]{40,}={0,2}(?![A-Za-z0-9+/])")
SAYI = re.compile(r"^-?(0|[1-9]\d*)([.,]\d+)?$")


def oku(yol):
    uz = os.path.splitext(yol)[1].lower()
    if uz == ".json":
        with open(yol, encoding="utf-8-sig") as f:
            veri = json.load(f)
        if isinstance(veri, dict):
            veri = veri.get("satirlar", veri.get("rows"))
        if not isinstance(veri, list) or not all(isinstance(x, dict) for x in veri):
            sys.exit(f"HATA: {yol}: JSON nesne listesi ya da {{'satirlar': [...]}} olmalı")
        kol = []
        for x in veri:
            for k in x:
                if k not in kol:
                    kol.append(k)
        return kol, [[x.get(k) for k in kol] for x in veri]
    with open(yol, encoding="utf-8-sig", newline="") as f:
        metin = f.read()
    try:
        lehce = csv.Sniffer().sniff(metin[:20000], delimiters=",;|\t")
    except csv.Error:
        lehce = csv.excel
    satirlar = list(csv.reader(io.StringIO(metin, newline=""), lehce))
    if not satirlar:
        sys.exit(f"HATA: {yol} boş")
    kol, govde = satirlar[0], [r for r in satirlar[1:] if any(c.strip() for c in r)]
    for i, r in enumerate(govde, 2):
        if len(r) != len(kol):
            sys.exit(f"HATA: {yol} satır {i}: {len(r)} alan, başlıkta {len(kol)}")
    return kol, govde


def sutun_sec(kol, govde, spek):
    if not spek:
        return kol, govde
    secim = []
    for parca in [p.strip() for p in spek.split(",") if p.strip()]:
        kaynak, _, yeni = parca.partition("=")
        if kaynak not in kol:
            sys.exit(f"HATA: sütun yok: {kaynak!r}. Var olanlar: {kol}")
        secim.append((kol.index(kaynak), yeni or kaynak))
    return [y for _, y in secim], [[r[i] for i, _ in secim] for r in govde]


def deger(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, (list, dict)):
        return json.dumps(v, ensure_ascii=False)
    s = str(v).strip()
    # baştaki sıfır korunur (ör. "007" kod), yalnız saf sayı dönüşür
    if SAYI.match(s) and not (len(s) > 1 and s.startswith("0") and s[1:2] not in ".,"):
        try:
            return int(s) if re.fullmatch(r"-?\d+", s) else float(s.replace(",", "."))
        except ValueError:
            return s
    return s


def sayfa_adi(ad, kullanilan):
    ad = re.sub(r"[\[\]:*?/\\]", "_", ad).strip() or "Veri"
    ad = ad[:31]
    temel, i = ad, 2
    while ad.lower() in kullanilan:
        ad = f"{temel[:28]}_{i}"
        i += 1
    kullanilan.add(ad.lower())
    return ad


def yaz_hucre(ws, r, c, v):
    h = ws.cell(r, c, v)
    if isinstance(v, str) and v.startswith("="):
        h.data_type = TYPE_STRING  # formül olarak yorumlanmasın
    return h


def veri_sayfasi(wb, ad, kol, govde, dondur_sutun):
    ws = wb.create_sheet(ad)
    for c, k in enumerate(kol, 1):
        h = yaz_hucre(ws, 1, c, str(k))
        h.font, h.fill, h.border = BASLIK_FONT, BASLIK_FILL, KENAR
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    genis = [len(str(k)) for k in kol]
    for r, satir in enumerate(govde, 2):
        for c, v in enumerate(satir, 1):
            h = yaz_hucre(ws, r, c, deger(v))
            h.border = KENAR
            h.alignment = Alignment(vertical="top", wrap_text=True)
            uz = max((len(x) for x in str(h.value or "").split("\n")), default=0)
            genis[c - 1] = max(genis[c - 1], uz)
    for c, g in enumerate(genis, 1):
        ws.column_dimensions[get_column_letter(c)].width = max(MIN_GEN, min(MAX_GEN, g + 2))
    ws.freeze_panes = f"{get_column_letter(1 + dondur_sutun)}2"
    if kol:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(kol))}{max(1, len(govde) + 1)}"
    return ws


def ozet_sayfasi(wb, baslik, alt, ozet_yolu, sayfalar, kaynaklar):
    ws = wb.active
    ws.title = "Ozet"
    satirlar = [("baslik", baslik), ("alt", alt or dt.date.today().strftime("%d.%m.%Y")), ("bos", "")]
    if ozet_yolu:
        with open(ozet_yolu, encoding="utf-8-sig") as f:
            for l in f.read().splitlines():
                l = l.rstrip()
                if l.startswith("#"):
                    satirlar.append(("bolum", l.lstrip("#").strip().upper()))
                elif not l.strip():
                    satirlar.append(("bos", ""))
                else:
                    satirlar.append(("madde", re.sub(r"^\s*[-*]\s+", "", l)))
        if satirlar[-1][0] != "bos":
            satirlar.append(("bos", ""))
    satirlar.append(("bolum", "SAYFALAR"))
    for ad, n, aciklama in sayfalar:
        satirlar.append(("madde", f"{ad}: {n} satır" + (f". {aciklama}" if aciklama else "")))
    satirlar += [("bos", ""), ("bolum", "ÜRETİM")]
    satirlar.append(("madde", f"{dt.date.today().strftime('%d.%m.%Y')} · kaynak: {', '.join(kaynaklar)} · "
                              f"betik: orkestra sade-excel skill'i (sade_excel.py)"))
    for r, (tur, metin) in enumerate(satirlar, 1):
        h = yaz_hucre(ws, r, 1, metin or None)
        h.alignment = Alignment(wrap_text=True, vertical="top")
        if tur == "baslik":
            h.font = Font(bold=True, size=16, color=KOYU)
        elif tur == "alt":
            h.font = Font(italic=True, color=GRI)
        elif tur == "bolum":
            h.font, h.fill = BASLIK_FONT, BASLIK_FILL
    ws.column_dimensions["A"].width = 150
    return len(satirlar)


def sir_tara(wb):
    kesin, suphe = [], []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for h in row:
                if isinstance(h.value, str):
                    if any(p.search(h.value) for p in SIR_KESIN):
                        kesin.append(f"{ws.title}!{h.coordinate}")
                    elif SIR_SUPHE.search(h.value):
                        suphe.append(f"{ws.title}!{h.coordinate}")
    return kesin, suphe


def masaustu():
    if os.name == "nt":
        try:
            import ctypes
            from ctypes import wintypes
            buf = ctypes.create_unicode_buffer(wintypes.MAX_PATH)
            if ctypes.windll.shell32.SHGetFolderPathW(None, 0x0010, None, 0, buf) == 0 and buf.value:
                return buf.value  # CSIDL_DESKTOPDIRECTORY: gerçek masaüstü
        except Exception:
            pass
    return os.path.join(os.path.expanduser("~"), "Desktop")


def dogrula(yol, beklenen):
    wb = load_workbook(yol)
    hatalar = []
    if wb.sheetnames != [a for a, _, _ in beklenen]:
        hatalar.append(f"sayfa adları {wb.sheetnames} != {[a for a, _, _ in beklenen]}")
    for ad, n_satir, n_kol in beklenen:
        if ad not in wb.sheetnames:
            continue
        ws = wb[ad]
        if ws.max_row != n_satir or (n_kol and ws.max_column != n_kol):
            hatalar.append(f"{ad}: {ws.max_row}x{ws.max_column}, beklenen {n_satir}x{n_kol}")
        if ad != "Ozet" and (ws.freeze_panes is None or ws.auto_filter.ref is None):
            hatalar.append(f"{ad}: donmuş başlık ya da filtre yok")
    return hatalar


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--cikti", required=True)
    ap.add_argument("--baslik", required=True)
    ap.add_argument("--alt-baslik")
    ap.add_argument("--veri", action="append", required=True)
    ap.add_argument("--sayfa", action="append", default=[])
    ap.add_argument("--sutunlar", action="append", default=[])
    ap.add_argument("--aciklama", action="append", default=[])
    ap.add_argument("--ozet")
    ap.add_argument("--dondur-sutun", type=int, default=0, help="başlıkla birlikte dondurulacak sol sütun sayısı")
    ap.add_argument("--masaustu", action="store_true", help="çıktının kopyasını masaüstüne koy")
    a = ap.parse_args()

    if len(a.veri) > MAKS_VERI_SAYFASI:
        sys.exit(f"HATA: en çok {MAKS_VERI_SAYFASI} veri sayfası (sade kuralı); {len(a.veri)} verildi. "
                 "Veriyi birleştir ya da ayrı dosya üret.")
    if not a.cikti.lower().endswith(".xlsx"):
        sys.exit("HATA: --cikti .xlsx ile bitmeli")
    for y in a.veri + ([a.ozet] if a.ozet else []):
        if not os.path.isfile(y):
            sys.exit(f"HATA: dosya yok: {y}")

    wb = Workbook()
    kullanilan = {"ozet"}
    tablolar = []
    for i, yol in enumerate(a.veri):
        kol, govde = oku(yol)
        kol, govde = sutun_sec(kol, govde, a.sutunlar[i] if i < len(a.sutunlar) else "")
        ad = sayfa_adi(a.sayfa[i] if i < len(a.sayfa) else os.path.splitext(os.path.basename(yol))[0], kullanilan)
        tablolar.append((ad, kol, govde, a.aciklama[i] if i < len(a.aciklama) else ""))
    n_ozet = ozet_sayfasi(wb, a.baslik, a.alt_baslik, a.ozet,
                          [(ad, len(g), ac) for ad, _, g, ac in tablolar],
                          [os.path.basename(y) for y in a.veri])
    for ad, kol, govde, _ in tablolar:
        veri_sayfasi(wb, ad, kol, govde, a.dondur_sutun)

    kesin, suphe = sir_tara(wb)
    if kesin:
        print("HATA: sır kalıbı bulundu, dosya yazılmadı. Hücreler:", ", ".join(kesin[:20]))
        sys.exit(3)
    if suphe:
        print("UYARI: uzun anahtar benzeri dizi (elle bak):", ", ".join(suphe[:20]))

    os.makedirs(os.path.dirname(os.path.abspath(a.cikti)), exist_ok=True)
    wb.save(a.cikti)
    beklenen = [("Ozet", n_ozet, 1)] + [(ad, len(g) + 1, len(k)) for ad, k, g, _ in tablolar]
    hatalar = dogrula(a.cikti, beklenen)
    for ad, k, g, _ in tablolar:
        print(f"  {ad}: {len(g)} satır, {len(k)} sütun")
    if hatalar:
        print("TUTARSIZ:", "; ".join(hatalar))
        sys.exit(2)
    print("TUTARLI ->", os.path.abspath(a.cikti))
    if a.masaustu:
        hedef = os.path.join(masaustu(), os.path.basename(a.cikti))
        shutil.copy2(a.cikti, hedef)
        print("Masaüstü kopyası ->", hedef)


if __name__ == "__main__":
    main()
